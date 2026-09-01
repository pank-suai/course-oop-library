#!/usr/bin/env python3
"""Build the final Russian 16:9 defense deck from deck.json."""
from __future__ import annotations

import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
PROJECT_ROOT = ROOT.parent
DATA = json.loads((ROOT / "deck.json").read_text(encoding="utf-8"))
OUT = ROOT / "course-oop-library.pptx"


def verified_test_summary():
    """Read the latest verified count instead of baking it into the deck."""
    evidence = (PROJECT_ROOT / "artifacts" / "test-results.txt").read_text(encoding="utf-8")
    match = re.search(
        r"(?:contain\s+)?\*{0,2}(\d+) tests,\s*(\d+) failures,\s*(\d+) errors,\s*(?:and\s+)?(\d+) skipped",
        evidence,
    )
    if not match:
        raise RuntimeError("final test count was not found in artifacts/test-results.txt")
    total, failures, errors, skipped = map(int, match.groups())
    if failures or errors or skipped:
        raise RuntimeError("final test evidence is not a clean pass")
    return total

W, H = 13.333, 7.5
BG = "F8FAFC"
NAVY = "17324D"
INK = "172A3A"
MUTED = "476275"
BLUE = "E9F4FB"
LAVENDER = "F0EDFA"
GREEN = "EDF8F2"
GOLD = "FFF5DF"
ROSE = "FFF0EF"
BORDER = "B6CBD5"
TEAL = "24739B"
PURPLE = "735A9B"
GREEN_LINE = "328366"
GOLD_LINE = "B87517"

COLORS = {name: RGBColor.from_string(value) for name, value in {
    "bg": BG, "navy": NAVY, "ink": INK, "muted": MUTED, "blue": BLUE,
    "lavender": LAVENDER, "green": GREEN, "gold": GOLD, "rose": ROSE,
    "border": BORDER, "teal": TEAL, "purple": PURPLE, "green_line": GREEN_LINE,
    "gold_line": GOLD_LINE,
}.items()}


def rgb(value):
    return COLORS[value] if value in COLORS else RGBColor.from_string(value)


def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLORS["bg"]


def add_box(slide, x, y, w, h, fill=BG, line=BORDER, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(1.2)
    return shape


def add_text(slide, text, x, y, w, h, size=28, color="ink", bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.08,
             font="Arial"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.bold = bold
    return shape


def add_rich_card(slide, heading, body, x, y, w, h, fill="blue", line="teal"):
    add_box(slide, x, y, w, h, fill, line)
    shape = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.10), Inches(w - 0.36), Inches(h - 0.20))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run(); r.text = heading
    r.font.name = "Arial"; r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = COLORS["navy"]
    p2 = tf.add_paragraph(); p2.space_before = Pt(0)
    r2 = p2.add_run(); r2.text = body
    r2.font.name = "Arial"; r2.font.size = Pt(28); r2.font.color.rgb = COLORS["ink"]
    return shape


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.56, 0.22, 12.2, 0.58, size=36, color="navy", bold=True, valign=MSO_ANCHOR.MIDDLE)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.56), Inches(0.91), Inches(12.2), Inches(0.025))
    line.fill.solid(); line.fill.fore_color.rgb = COLORS["border"]; line.line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 0.58, 0.98, 12.1, 0.36, size=28, color="muted")


def add_image(slide, name, x, y, w, h):
    slide.shapes.add_picture(str(ROOT / "assets" / name), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def slide_title(slide, spec):
    set_bg(slide)
    add_title(slide, spec["title"])


def build_title(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
    # restrained academic accent block
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.32), Inches(H))
    accent.fill.solid(); accent.fill.fore_color.rgb = COLORS["teal"]; accent.line.fill.background()
    add_text(slide, "КУРСОВОЙ ПРОЕКТ", 0.80, 0.74, 5.2, 0.42, size=28, color="teal", bold=True)
    add_text(slide, spec["title"], 0.78, 1.55, 11.7, 1.55, size=44, color="navy", bold=True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, spec["subtitle"], 0.82, 3.55, 11.15, 0.42, size=28, color="muted")
    add_text(slide, spec["author"], 0.82, 3.98, 11.15, 0.42, size=28, color="muted")
    add_box(slide, 0.82, 5.55, 5.15, 0.82, "blue", "teal")
    add_text(slide, "Room-backed SQLite · Compose Desktop", 1.00, 5.67, 4.80, 0.48, size=28, color="navy", bold=True)
    add_text(slide, "Защита проекта", 9.35, 5.72, 3.0, 0.46, size=30, color="teal", bold=True, align=PP_ALIGN.RIGHT)


def build_cards(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide_title(slide, spec)
    cards = spec["cards"]
    positions = [(0.72, 1.42, 5.85, 1.55), (6.78, 1.42, 5.85, 1.55), (0.72, 3.38, 11.91, 1.48)]
    fills = [("rose", "gold_line"), ("blue", "teal"), ("lavender", "purple")]
    for card, (x, y, w, h), (fill, line) in zip(cards, positions, fills):
        add_rich_card(slide, card["heading"], card["text"], x, y, w, h, fill, line)
    add_text(slide, "Предметная область → доменная модель → сохранение → пользовательский сценарий", 0.82, 5.62, 11.7, 0.60, size=30, color="navy", bold=True, align=PP_ALIGN.CENTER)


def build_requirements(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide_title(slide, spec)
    positions = [(0.70, 1.34), (6.86, 1.34), (0.70, 2.92), (6.86, 2.92), (0.70, 4.50)]
    fills = [("blue", "teal"), ("gold", "gold_line"), ("green", "green_line"), ("lavender", "purple"), ("rose", "gold_line")]
    for i, (text, (x, y), (fill, line)) in enumerate(zip(spec["items"], positions, fills), 1):
        w = 5.78 if i != 5 else 11.94
        add_box(slide, x, y, w, 1.22, fill, line)
        add_text(slide, f"{i}.  {text}", x + 0.18, y + 0.07, w - 0.36, 1.08, size=28, color="ink", bold=(i == 1))


def build_image(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide_title(slide, spec)
    add_image(slide, spec["image"], 0.42, 1.14, 12.49, 6.05)


def build_patterns_overview(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide_title(slide, spec)
    positions = [(0.70, 1.30), (6.86, 1.30), (0.70, 2.70), (6.86, 2.70), (0.70, 4.10), (6.86, 4.10), (0.70, 5.50)]
    fills = [("blue", "teal"), ("lavender", "purple"), ("gold", "gold_line"), ("green", "green_line"), ("blue", "teal"), ("rose", "gold_line"), ("lavender", "purple")]
    for (heading, body), (x, y), (fill, line) in zip(spec["items"], positions, fills):
        w = 5.78
        add_rich_card(slide, heading, body, x, y, w, 1.10, fill, line)


def build_dual_image(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide_title(slide, spec)
    add_image(slide, spec["images"][0], 0.46, 1.20, 6.12, 3.44)
    add_image(slide, spec["images"][1], 6.76, 1.20, 6.12, 3.44)
    add_box(slide, 0.78, 5.02, 11.78, 1.10, "blue", "teal")
    add_text(slide, spec["caption"], 1.00, 5.17, 11.34, 0.78, size=28, color="navy", bold=True, align=PP_ALIGN.CENTER)


def build_triple_image(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide_title(slide, spec)
    xs = [0.30, 4.61, 8.92]
    for x, image in zip(xs, spec["images"]):
        add_image(slide, image, x, 1.18, 4.12, 5.98)


def build_transactions(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide_title(slide, spec)
    positions = [(0.64, 1.36), (4.56, 1.36), (8.48, 1.36)]
    fills = [("blue", "teal"), ("green", "green_line"), ("gold", "gold_line")]
    for card, (x, y), (fill, line) in zip(spec["cards"], positions, fills):
        add_rich_card(slide, card["heading"], card["text"], x, y, 3.72, 3.55, fill, line)
    add_text(slide, "Атомарность защищает инвентарь и связи между сущностями", 0.88, 5.55, 11.55, 0.68, size=30, color="navy", bold=True, align=PP_ALIGN.CENTER)


def build_flows(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide_title(slide, spec)
    positions = [(0.70, 1.38), (6.86, 1.38), (0.70, 3.18), (6.86, 3.18)]
    fills = [("blue", "teal"), ("green", "green_line"), ("gold", "gold_line"), ("lavender", "purple")]
    for (heading, body), (x, y), (fill, line) in zip(spec["flows"], positions, fills):
        add_rich_card(slide, heading, body, x, y, 5.78, 1.44, fill, line)
    add_box(slide, 0.78, 5.40, 11.76, 0.90, "rose", "gold_line")
    add_text(slide, spec["caption"], 1.00, 5.50, 11.30, 0.70, size=28, color="navy", bold=True, align=PP_ALIGN.CENTER)


def build_tests(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide_title(slide, spec)
    total = verified_test_summary()
    if sum(int(row[2]) for row in spec["rows"]) != total:
        raise RuntimeError("deck test rows do not match the latest verified total")
    headline = f"{total} тестов · 0 ошибок · 0 пропусков"
    add_box(slide, 0.72, 1.20, 11.90, 0.82, "green", "green_line")
    add_text(slide, headline, 0.94, 1.31, 11.46, 0.58, size=36, color="green_line", bold=True, align=PP_ALIGN.CENTER)
    # Compact evidence list; every visible run remains at least 28 pt.
    row_y = 2.23
    for i, row in enumerate(spec["rows"]):
        fill, line = (("blue", "teal") if i % 2 == 0 else ("bg", "border"))
        add_box(slide, 0.72, row_y + i * 0.57, 7.28, 0.48, fill, line, radius=False)
        add_text(slide, f"{row[0]}  ·  {row[1]}  ·  {row[2]}  ·  {row[3]}", 0.90, row_y + i * 0.57 + 0.02, 6.90, 0.40, size=28, color="ink", bold=(row[3] == "PASS"))
    add_box(slide, 8.25, 2.23, 4.38, 3.40, "gold", "gold_line")
    add_text(slide, spec["control"], 8.52, 2.50, 3.84, 2.80, size=28, color="ink", bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Источник: clean test build и JUnit XML", 0.84, 6.10, 7.10, 0.42, size=28, color="muted")


def build_conclusion(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide_title(slide, spec)
    for i, text in enumerate(spec["items"]):
        y = 1.34 + i * 0.96
        add_box(slide, 0.78, y, 11.78, 0.76, "blue" if i % 2 == 0 else "green", "teal" if i % 2 == 0 else "green_line")
        add_text(slide, f"✓  {text}", 1.02, y + 0.06, 11.30, 0.62, size=28, color="ink", bold=(i == 4))
    add_text(slide, "Спасибо за внимание", 0.80, 6.34, 11.72, 0.52, size=34, color="navy", bold=True, align=PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    builders = {
        "title": build_title,
        "cards": build_cards,
        "requirements": build_requirements,
        "image": build_image,
        "patterns_overview": build_patterns_overview,
        "dual_image": build_dual_image,
        "triple_image": build_triple_image,
        "transactions": build_transactions,
        "flows": build_flows,
        "tests": build_tests,
        "conclusion": build_conclusion,
    }
    for spec in DATA["slides"]:
        builders[spec["kind"]](prs, spec)
    prs.save(OUT)
    print(json.dumps({"ok": True, "output": str(OUT), "slides": len(prs.slides)}, ensure_ascii=False))


if __name__ == "__main__":
    build()
