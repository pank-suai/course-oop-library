#!/usr/bin/env python3
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt

path = Path(__file__).with_name("course-oop-library.pptx")
prs = Presentation(path)
minimum = Pt(28)
for slide in prs.slides:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        shape.text_frame.word_wrap = True
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size is None or run.font.size < minimum:
                    run.font.size = minimum
                if run.font.name is None:
                    run.font.name = "Arial"
prs.save(path)

# Read back and fail if any visible run is below 28 pt.
verified = Presentation(path)
violations = []
for slide_index, slide in enumerate(verified.slides, start=1):
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text.strip() and (run.font.size is None or run.font.size < minimum):
                    violations.append((slide_index, run.text, run.font.size.pt if run.font.size else None))
if violations:
    raise SystemExit(f"font-size violations: {violations}")
print(f"verified: {len(verified.slides)} slides, all visible runs >= 28 pt")
